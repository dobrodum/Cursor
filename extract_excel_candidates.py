#!/usr/bin/env python3
"""Extract empirical and regression candidate rows from Excel model files."""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import xlwings as xw
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# -----------------------------------------------------------------------------
# Configure these two paths before running.
# -----------------------------------------------------------------------------
input_dir = Path("/path/to/input")
output_dir = Path("/path/to/output")

EMPIRICAL_OUTPUT_COLUMNS = [
    "model",
    "ticker",
    "model_period",
    "model_date",
    "method",
    "parameter_name",
    "parameter_value",
    "num_quarters_used",
    "last_quarter_used",
    "forecast_value",
    "actual_value",
    "forecast_max",
    "forecast_min",
    "range_width",
    "avg_penetration_pct",
    "quarterly_sales",
    "reported_sales",
    "growth_rate_pct",
    "sales_captured_in_db_pct",
    "source_file",
]

REGRESSION_OUTPUT_COLUMNS = [
    "model",
    "ticker",
    "model_period",
    "model_date",
    "method",
    "parameter_name",
    "parameter_value",
    "num_quarters_used",
    "forecast_value",
    "actual_value",
    "forecast_max",
    "forecast_min",
    "range_width",
    "intercept",
    "slope",
    "source_file",
]

N_QUARTERS = 10
EMPIRICAL_SHEET_NAME = "Empirical Model"
REGRESSION_SHEET_NAME = "Regression Model"
EPSILON = 1e-9


@dataclass
class SheetCache:
    sheet: xw.Sheet
    values: List[List[Any]]
    top_row: int
    left_col: int
    bottom_row: int
    right_col: int
    labels: Dict[str, List[Tuple[int, int]]]


def normalize_2d(values: Any) -> List[List[Any]]:
    if values is None:
        return []
    if not isinstance(values, list):
        return [[values]]
    if values and not isinstance(values[0], list):
        return [values]
    return values


def normalize_label(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    normalized = re.sub(r"\s+", " ", value.strip().lower())
    normalized = normalized.replace("_", " ")
    return normalized


def to_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        text = text.replace(",", "")
        text = text.replace("$", "")
        text = text.replace("%", "")
        try:
            return float(text)
        except ValueError:
            return None
    return None


def maybe_round(value: Optional[float], digits: int = 10) -> Optional[float]:
    if value is None:
        return None
    return round(value, digits)


def build_sheet_cache(sheet: xw.Sheet) -> SheetCache:
    used_range = sheet.used_range
    values = normalize_2d(used_range.value)
    top_row = used_range.row
    left_col = used_range.column
    row_count = len(values)
    col_count = max((len(row) for row in values), default=0)
    bottom_row = top_row + row_count - 1 if row_count else top_row
    right_col = left_col + col_count - 1 if col_count else left_col

    labels: Dict[str, List[Tuple[int, int]]] = {}
    for r_idx, row in enumerate(values):
        for c_idx, value in enumerate(row):
            key = normalize_label(value)
            if key:
                labels.setdefault(key, []).append((top_row + r_idx, left_col + c_idx))

    return SheetCache(
        sheet=sheet,
        values=values,
        top_row=top_row,
        left_col=left_col,
        bottom_row=bottom_row,
        right_col=right_col,
        labels=labels,
    )


def cache_value(cache: SheetCache, row: int, col: int) -> Any:
    if row < cache.top_row or col < cache.left_col:
        return None
    r_idx = row - cache.top_row
    c_idx = col - cache.left_col
    if r_idx >= len(cache.values) or r_idx < 0:
        return None
    row_values = cache.values[r_idx]
    if c_idx >= len(row_values) or c_idx < 0:
        return None
    return row_values[c_idx]


def find_first_label(
    cache: SheetCache, exact_candidates: Sequence[str], contains_candidates: Sequence[str] = ()
) -> Optional[Tuple[int, int]]:
    for candidate in exact_candidates:
        key = normalize_label(candidate)
        hits = cache.labels.get(key)
        if hits:
            return hits[0]
    for key, positions in cache.labels.items():
        for token in contains_candidates:
            if token in key:
                return positions[0]
    return None


def find_anchor_max(cache: SheetCache) -> Optional[Tuple[int, int]]:
    return find_first_label(cache, exact_candidates=("max",), contains_candidates=(" max", "max "))


def nearest_numeric_to_right(
    cache: SheetCache, row: int, col: int, max_steps: int = 8
) -> Optional[float]:
    for step in range(1, max_steps + 1):
        value = to_float(cache_value(cache, row, col + step))
        if value is not None:
            return value
    return None


def value_from_label_right(
    cache: SheetCache,
    exact_labels: Sequence[str],
    contains_labels: Sequence[str] = (),
    max_steps: int = 8,
) -> Optional[float]:
    hit = find_first_label(cache, exact_labels, contains_labels)
    if not hit:
        return None
    row, col = hit
    return nearest_numeric_to_right(cache, row, col, max_steps=max_steps)


def row_for_label(
    cache: SheetCache, exact_labels: Sequence[str], contains_labels: Sequence[str] = ()
) -> Optional[int]:
    hit = find_first_label(cache, exact_labels, contains_labels)
    return hit[0] if hit else None


def row_for_label_near_anchor(
    cache: SheetCache,
    anchor_row: int,
    exact_labels: Sequence[str],
    contains_labels: Sequence[str] = (),
    max_distance: int = 40,
) -> Optional[int]:
    candidates: List[int] = []
    for candidate in exact_labels:
        key = normalize_label(candidate)
        candidates.extend(row for row, _ in cache.labels.get(key, []))

    for key, positions in cache.labels.items():
        if any(token in key for token in contains_labels):
            candidates.extend(row for row, _ in positions)

    if not candidates:
        return None

    nearest = min(candidates, key=lambda row: abs(row - anchor_row))
    if abs(nearest - anchor_row) > max_distance:
        return None
    return nearest


def numeric_columns_for_rows(
    cache: SheetCache, row_a: int, row_b: int, right_bound_exclusive: int
) -> List[int]:
    cols: List[int] = []
    for col in range(cache.left_col, right_bound_exclusive):
        a = to_float(cache_value(cache, row_a, col))
        b = to_float(cache_value(cache, row_b, col))
        if a is None or b is None:
            continue
        if abs(a) <= EPSILON:
            continue
        cols.append(col)
    return cols


def first_non_empty_in_row(cache: SheetCache, row: int, col: int, max_steps: int = 8) -> Any:
    for step in range(1, max_steps + 1):
        value = cache_value(cache, row, col + step)
        if value not in (None, ""):
            return value
    return None


def parse_file_labels(file_path: Path) -> Dict[str, str]:
    stem = file_path.stem
    parts = [part.strip() for part in stem.split(" - ")]

    ticker = ""
    if len(parts) >= 2:
        ticker = parts[1]
    else:
        upper_tokens = re.findall(r"\b[A-Z]{2,6}\b", stem)
        ticker = upper_tokens[0] if upper_tokens else ""

    period_match = re.search(
        r"(?i)\b(early|mid|late)\s*([A-Za-z]{3,9})\s*(20\d{2})\b", stem
    )
    model_period = ""
    model_date = ""
    if period_match:
        phase_raw, month_raw, year_raw = period_match.groups()
        phase = phase_raw.title()
        month_token = month_raw.title()
        year = int(year_raw)
        month_number = parse_month(month_token)
        if month_number is not None:
            day = {"Early": 5, "Mid": 15, "Late": 25}[phase]
            model_period = f"{phase}{dt.date(year, month_number, 1):%b}_{year}"
            model_date = dt.date(year, month_number, day).isoformat()

    if not model_period:
        model_period = "unknown_period"
    if not model_date:
        model_date = ""
    if not ticker:
        ticker = "unknown_ticker"

    model = f"{ticker}_{model_period}"
    return {
        "ticker": ticker,
        "model_period": model_period,
        "model_date": model_date,
        "model": model,
    }


def parse_month(month_token: str) -> Optional[int]:
    for fmt in ("%b", "%B"):
        try:
            return dt.datetime.strptime(month_token, fmt).month
        except ValueError:
            continue
    return None


def safe_close_workbook(wb: Optional[xw.Book]) -> None:
    if wb is None:
        return
    try:
        wb.close(save=False)
        return
    except TypeError:
        pass
    except Exception:
        pass

    try:
        wb.close(False)
        return
    except Exception:
        pass

    try:
        wb.api.Close(SaveChanges=False)
    except Exception as exc:
        print(f"warning: unable to close workbook safely: {exc}")


def find_forecast_bounds(cache: SheetCache, anchor: Tuple[int, int]) -> Tuple[Optional[float], Optional[float]]:
    anchor_row, anchor_col = anchor
    max_val = nearest_numeric_to_right(cache, anchor_row, anchor_col)
    min_val = value_from_label_right(cache, exact_labels=("min",), contains_labels=(" min", "min "))
    return max_val, min_val


def build_empirical_rows(
    wb: xw.Book, cache: SheetCache, file_meta: Dict[str, str], source_name: str
) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    anchor = find_anchor_max(cache)
    if not anchor:
        print(f"  skipped empirical sheet: no 'max' anchor found")
        return rows

    anchor_row, anchor_col = anchor
    max_val, min_val = find_forecast_bounds(cache, anchor)
    range_width = (max_val - min_val) if (max_val is not None and min_val is not None) else None

    quarterly_row = row_for_label_near_anchor(
        cache,
        anchor_row=anchor_row,
        exact_labels=("quarterly sales", "quarterly_sales"),
        contains_labels=("quarterly sale",),
    )
    reported_row = row_for_label_near_anchor(
        cache,
        anchor_row=anchor_row,
        exact_labels=("reported sales", "reported_sales"),
        contains_labels=("reported sale",),
    )
    growth_row = row_for_label_near_anchor(
        cache,
        anchor_row=anchor_row,
        exact_labels=("growth rate", "growth rate %", "growth_rate_pct"),
        contains_labels=("growth rate",),
    )
    captured_row = row_for_label_near_anchor(
        cache,
        anchor_row=anchor_row,
        exact_labels=(
            "sales captured in db %",
            "sales captured in db pct",
            "sales_captured_in_db_pct",
        ),
        contains_labels=("captured in db",),
    )

    if quarterly_row is None or reported_row is None:
        print("  skipped empirical sheet: missing quarterly/reported sales rows")
        return rows

    history_cols = numeric_columns_for_rows(
        cache, row_a=quarterly_row, row_b=reported_row, right_bound_exclusive=anchor_col
    )
    if not history_cols:
        print("  skipped empirical sheet: no overlapping numeric quarter history")
        return rows

    max_quarters = min(N_QUARTERS, len(history_cols))
    scratch_col = max(cache.right_col, anchor_col) + 20
    scratch_row = max(anchor_row, cache.top_row + 2)
    avg_cell = cache.sheet.cells(scratch_row, scratch_col)

    quarter_label_row = quarterly_row - 1 if quarterly_row > cache.top_row else quarterly_row
    for num_quarters in range(1, max_quarters + 1):
        window_cols = history_cols[-num_quarters:]
        start_col = window_cols[0]
        end_col = window_cols[-1]

        avg_formula = (
            f'=IFERROR(AVERAGE(R{reported_row}C{start_col}:R{reported_row}C{end_col}/'
            f'R{quarterly_row}C{start_col}:R{quarterly_row}C{end_col}),"")'
        )
        avg_cell.formula2 = avg_formula
        wb.app.calculate()
        avg_penetration_pct = to_float(avg_cell.value)

        quarter_header = cache_value(cache, quarter_label_row, start_col)
        last_quarter_used = str(quarter_header) if quarter_header not in (None, "") else ""

        latest_col = end_col
        quarterly_sales = to_float(cache_value(cache, quarterly_row, latest_col))
        reported_sales = to_float(cache_value(cache, reported_row, latest_col))
        growth_rate_pct = to_float(cache_value(cache, growth_row, latest_col)) if growth_row else None
        sales_captured_in_db_pct = (
            to_float(cache_value(cache, captured_row, latest_col)) if captured_row else None
        )

        if sales_captured_in_db_pct is None and quarterly_sales and reported_sales and abs(reported_sales) > EPSILON:
            sales_captured_in_db_pct = quarterly_sales / reported_sales

        forecast_value = None
        if quarterly_sales is not None and avg_penetration_pct not in (None, 0):
            forecast_value = quarterly_sales / avg_penetration_pct

        actual_value = reported_sales

        row = {
            "model": file_meta["model"],
            "ticker": file_meta["ticker"],
            "model_period": file_meta["model_period"],
            "model_date": file_meta["model_date"],
            "method": "empirical",
            "parameter_name": "avg_penetration_pct",
            "parameter_value": maybe_round(avg_penetration_pct),
            "num_quarters_used": num_quarters,
            "last_quarter_used": last_quarter_used,
            "forecast_value": maybe_round(forecast_value),
            "actual_value": maybe_round(actual_value),
            "forecast_max": maybe_round(max_val),
            "forecast_min": maybe_round(min_val),
            "range_width": maybe_round(range_width),
            "avg_penetration_pct": maybe_round(avg_penetration_pct),
            "quarterly_sales": maybe_round(quarterly_sales),
            "reported_sales": maybe_round(reported_sales),
            "growth_rate_pct": maybe_round(growth_rate_pct),
            "sales_captured_in_db_pct": maybe_round(sales_captured_in_db_pct),
            "source_file": source_name,
        }
        rows.append(row)

    return rows


def regression_data_rows(
    cache: SheetCache, x_col: int, y_col: int, anchor_row: int
) -> List[int]:
    rows: List[int] = []
    for row in range(cache.top_row, anchor_row):
        x_val = to_float(cache_value(cache, row, x_col))
        y_val = to_float(cache_value(cache, row, y_col))
        if x_val is None or y_val is None:
            continue
        rows.append(row)
    return rows


def values_match_for_dedup(previous: Dict[str, Any], current: Dict[str, Any]) -> bool:
    keys = ("forecast_value", "forecast_max", "forecast_min", "intercept", "slope")
    for key in keys:
        prev_val = previous.get(key)
        cur_val = current.get(key)
        if prev_val is None and cur_val is None:
            continue
        if prev_val is None or cur_val is None:
            return False
        if abs(float(prev_val) - float(cur_val)) > EPSILON:
            return False
    return True


def build_regression_rows(
    wb: xw.Book, cache: SheetCache, file_meta: Dict[str, str], source_name: str
) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    anchor = find_anchor_max(cache)
    if not anchor:
        print("  skipped regression sheet: no 'max' anchor found")
        return rows

    anchor_row, anchor_col = anchor
    y_col = anchor_col - 7
    x_col = anchor_col - 11

    data_rows = regression_data_rows(cache, x_col=x_col, y_col=y_col, anchor_row=anchor_row)
    if len(data_rows) < 2:
        print("  skipped regression sheet: not enough regression data points")
        return rows

    max_val, min_val = find_forecast_bounds(cache, anchor)
    range_width = (max_val - min_val) if (max_val is not None and min_val is not None) else None

    max_quarters = min(N_QUARTERS, len(data_rows))
    scratch_col = max(cache.right_col, anchor_col) + 24
    intercept_cell = cache.sheet.cells(anchor_row, scratch_col)
    slope_cell = cache.sheet.cells(anchor_row + 1, scratch_col)
    actual_value = value_from_label_right(
        cache,
        exact_labels=("actual value", "actual", "reported sales"),
        contains_labels=("actual",),
    )

    previous_row: Optional[Dict[str, Any]] = None
    for num_quarters in range(2, max_quarters + 1):
        window_rows = data_rows[-num_quarters:]
        start_row = window_rows[0]
        end_row = window_rows[-1]

        intercept_cell.formula2 = (
            f'=IFERROR(INTERCEPT(R{start_row}C{y_col}:R{end_row}C{y_col},'
            f'R{start_row}C{x_col}:R{end_row}C{x_col}),"")'
        )
        slope_cell.formula2 = (
            f'=IFERROR(SLOPE(R{start_row}C{y_col}:R{end_row}C{y_col},'
            f'R{start_row}C{x_col}:R{end_row}C{x_col}),"")'
        )
        wb.app.calculate()

        intercept = to_float(intercept_cell.value)
        slope = to_float(slope_cell.value)

        next_x = to_float(cache_value(cache, end_row + 1, x_col))
        if next_x is None:
            next_x = to_float(cache_value(cache, end_row, x_col))
        forecast_total_without_sa = None
        if intercept is not None and slope is not None and next_x is not None:
            forecast_total_without_sa = intercept + (slope * next_x)

        row: Dict[str, Any] = {
            "model": file_meta["model"],
            "ticker": file_meta["ticker"],
            "model_period": file_meta["model_period"],
            "model_date": file_meta["model_date"],
            "method": "regression",
            "parameter_name": "num_quarters_used",
            "parameter_value": num_quarters,
            "num_quarters_used": num_quarters,
            "forecast_value": maybe_round(forecast_total_without_sa),
            "actual_value": maybe_round(actual_value),
            "forecast_max": maybe_round(max_val),
            "forecast_min": maybe_round(min_val),
            "range_width": maybe_round(range_width),
            "intercept": maybe_round(intercept),
            "slope": maybe_round(slope),
            "source_file": source_name,
        }

        if previous_row and values_match_for_dedup(previous_row, row):
            continue

        rows.append(row)
        previous_row = row

    return rows


def resolve_output_path(input_folder: Path, destination_folder: Path) -> Path:
    destination_folder.mkdir(parents=True, exist_ok=True)
    base_name = f"{input_folder.name}_PARAM"
    candidate = destination_folder / f"{base_name}.xlsx"
    if not candidate.exists():
        return candidate

    suffix = 1
    while True:
        candidate = destination_folder / f"{base_name}.{suffix}.xlsx"
        if not candidate.exists():
            return candidate
        suffix += 1


def write_sheet(
    ws: Any,
    columns: Sequence[str],
    rows: Sequence[Dict[str, Any]],
) -> None:
    ws.append(list(columns))
    for row_dict in rows:
        ws.append([row_dict.get(column) for column in columns])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    bold_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold_font

    for col_index, column_name in enumerate(columns, start=1):
        max_len = len(column_name)
        for row in rows:
            value = row.get(column_name)
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        ws.column_dimensions[get_column_letter(col_index)].width = min(max_len + 2, 45)


def write_output_workbook(
    output_path: Path,
    empirical_rows: Sequence[Dict[str, Any]],
    regression_rows: Sequence[Dict[str, Any]],
) -> None:
    wb = Workbook()
    empirical_ws = wb.active
    empirical_ws.title = "empirical_candidates"
    write_sheet(empirical_ws, EMPIRICAL_OUTPUT_COLUMNS, empirical_rows)

    regression_ws = wb.create_sheet("regression_candidates")
    write_sheet(regression_ws, REGRESSION_OUTPUT_COLUMNS, regression_rows)

    wb.save(output_path)


def iter_source_files(folder: Path) -> Iterable[Path]:
    output_pattern = re.compile(rf"^{re.escape(folder.name)}_PARAM(?:\.\d+)?\.xlsx$", re.IGNORECASE)
    for path in sorted(folder.iterdir()):
        if not path.is_file():
            continue
        if path.name.startswith("~"):
            print(f"skipped file: {path.name} (temporary file)")
            continue
        if path.suffix.lower() != ".xlsx":
            print(f"skipped file: {path.name} (not an .xlsx file)")
            continue
        if output_pattern.match(path.name):
            print(f"skipped file: {path.name} (existing extraction output)")
            continue
        yield path


def main() -> None:
    if not input_dir.exists() or not input_dir.is_dir():
        raise FileNotFoundError(f"input_dir does not exist or is not a directory: {input_dir}")

    output_path = resolve_output_path(input_dir, output_dir)

    empirical_rows: List[Dict[str, Any]] = []
    regression_rows: List[Dict[str, Any]] = []
    processed_files = 0

    app: Optional[xw.App] = None
    original_calculation_mode = None
    try:
        app = xw.App(visible=False, add_book=False)
        app.display_alerts = False
        app.screen_updating = False
        original_calculation_mode = app.calculation
        app.calculation = "manual"

        for file_path in iter_source_files(input_dir):
            wb: Optional[xw.Book] = None
            try:
                wb = app.books.open(str(file_path), update_links=False)
                file_meta = parse_file_labels(file_path)
                sheet_by_name = {sheet.name: sheet for sheet in wb.sheets}

                if EMPIRICAL_SHEET_NAME in sheet_by_name:
                    empirical_cache = build_sheet_cache(sheet_by_name[EMPIRICAL_SHEET_NAME])
                    empirical_rows.extend(
                        build_empirical_rows(
                            wb=wb,
                            cache=empirical_cache,
                            file_meta=file_meta,
                            source_name=file_path.name,
                        )
                    )
                else:
                    print(f"  skipped empirical sheet: '{EMPIRICAL_SHEET_NAME}' not found")

                if REGRESSION_SHEET_NAME in sheet_by_name:
                    regression_cache = build_sheet_cache(sheet_by_name[REGRESSION_SHEET_NAME])
                    regression_rows.extend(
                        build_regression_rows(
                            wb=wb,
                            cache=regression_cache,
                            file_meta=file_meta,
                            source_name=file_path.name,
                        )
                    )
                else:
                    print(f"  skipped regression sheet: '{REGRESSION_SHEET_NAME}' not found")

                processed_files += 1
                print(f"processed file: {file_path.name}")
            except Exception as exc:
                print(f"skipped file: {file_path.name} (error: {exc})")
            finally:
                safe_close_workbook(wb)
    finally:
        if app is not None:
            try:
                if original_calculation_mode is not None:
                    app.calculation = original_calculation_mode
            except Exception:
                pass
            app.quit()

    write_output_workbook(
        output_path=output_path,
        empirical_rows=empirical_rows,
        regression_rows=regression_rows,
    )

    print(f"output path: {output_path}")
    print(f"number of files processed: {processed_files}")
    print(f"number of empirical rows: {len(empirical_rows)}")
    print(f"number of regression rows: {len(regression_rows)}")


if __name__ == "__main__":
    main()
