#!/usr/bin/env python3
"""Extract empirical/regression candidate rows from model workbooks.

The script opens each source workbook exactly once, reads both model sheets while
it is open, and writes one consolidated output workbook with:
  - empirical_candidates
  - regression_candidates
"""

from __future__ import annotations

import calendar
import datetime as dt
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import xlwings as xw
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# -----------------------------
# User-configurable directories
# -----------------------------
input_dir = "/path/to/input"
output_dir = "/path/to/output"


EMPIRICAL_COLUMNS: List[str] = [
    "model",
    "ticker",
    "model_period",
    "model_date",
    "method",
    "parameter_name",
    "parameter_value",
    "num_quarters_used",
    "last_quarter_used",
    "forecast_value",
    "actual_value",
    "forecast_max",
    "forecast_min",
    "range_width",
    "avg_penetration_pct",
    "quarterly_sales",
    "reported_sales",
    "growth_rate_pct",
    "sales_captured_in_db_pct",
    "source_file",
]

REGRESSION_COLUMNS: List[str] = [
    "model",
    "ticker",
    "model_period",
    "model_date",
    "method",
    "parameter_name",
    "parameter_value",
    "num_quarters_used",
    "forecast_value",
    "actual_value",
    "forecast_max",
    "forecast_min",
    "range_width",
    "intercept",
    "slope",
    "source_file",
]

PHASE_DAY = {"early": 5, "mid": 15, "late": 25}
PERIOD_PATTERN = re.compile(r"(Early|Mid|Late)\s*([A-Za-z]{3,9})\s*(\d{4})", re.IGNORECASE)


def month_number(month_text: str) -> Optional[int]:
    """Parse month text to month number."""
    cleaned = month_text.strip()
    if not cleaned:
        return None
    key = cleaned[:3].title()
    for month_idx in range(1, 13):
        if calendar.month_abbr[month_idx] == key:
            return month_idx
    return None


def parse_file_labels(file_path: Path) -> Dict[str, str]:
    """Parse model metadata from a file name."""
    stem = file_path.stem
    parts = [part.strip() for part in stem.split(" - ")]
    ticker = parts[1].upper() if len(parts) >= 2 and parts[1] else "UNKNOWN"

    period_source = parts[2] if len(parts) >= 3 else ""
    period_source = period_source.split("_")[0]
    match = PERIOD_PATTERN.search(period_source)

    model_period = "unknown_period"
    model_date = ""

    if match:
        phase = match.group(1).title()
        month_text = match.group(2)
        year = int(match.group(3))
        month = month_number(month_text)
        if month:
            month_abbr = calendar.month_abbr[month]
            model_period = f"{phase}{month_abbr}_{year}"
            day = PHASE_DAY[phase.lower()]
            model_date = dt.date(year, month, day).isoformat()
    elif period_source:
        fallback = re.sub(r"[^A-Za-z0-9]+", "_", period_source).strip("_")
        if fallback:
            model_period = fallback

    model = f"{ticker}_{model_period}"
    return {
        "ticker": ticker,
        "model_period": model_period,
        "model_date": model_date,
        "model": model,
    }


def to_float(value: Any) -> Optional[float]:
    """Best-effort numeric conversion."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def ensure_2d(values: Any) -> List[List[Any]]:
    """Normalize xlwings used_range values to a 2D list."""
    if values is None:
        return []
    if not isinstance(values, list):
        return [[values]]
    if not values:
        return []
    if isinstance(values[0], list):
        return values  # type: ignore[return-value]
    return [values]  # type: ignore[return-value]


def find_anchor_cell(sheet: xw.Sheet, token: str = "max") -> Optional[Tuple[int, int]]:
    """Find the best matching anchor cell in one used_range scan."""
    used = sheet.used_range
    matrix = ensure_2d(used.value)
    if not matrix:
        return None

    top_row = used.row
    left_col = used.column
    token = token.lower()
    candidates: List[Tuple[int, int, int]] = []

    for row_idx, row in enumerate(matrix):
        for col_idx, value in enumerate(row):
            if not isinstance(value, str):
                continue
            text = value.strip().lower()
            if text == token:
                score = 3
            elif token in text:
                score = 1
            else:
                continue

            for delta in (-1, 1):
                near_idx = col_idx + delta
                if 0 <= near_idx < len(row):
                    near_val = row[near_idx]
                    if isinstance(near_val, str) and near_val.strip().lower().startswith("min"):
                        score += 1
                        break

            candidates.append((score, row_idx, col_idx))

    if not candidates:
        return None

    _, best_row_idx, best_col_idx = sorted(candidates, key=lambda item: (-item[0], item[1], item[2]))[0]
    return top_row + best_row_idx, left_col + best_col_idx


def set_formula2_r1c1(cell: xw.Range, formula_r1c1: str) -> None:
    """Set an R1C1 formula via Formula2 when available, with safe fallback."""
    try:
        cell.api.Formula2R1C1 = formula_r1c1
        return
    except Exception:
        pass
    try:
        cell.formula2 = formula_r1c1
        return
    except Exception:
        pass
    cell.api.FormulaR1C1 = formula_r1c1


def safe_close_workbook(wb: Optional[xw.Book]) -> None:
    """Close workbook without saving, with compatibility fallbacks."""
    if wb is None:
        return
    try:
        wb.close(save=False)
        return
    except TypeError:
        pass
    except Exception:
        pass

    for close_args in ((False,), tuple()):
        try:
            wb.close(*close_args)
            return
        except Exception:
            continue


def iterate_source_files(folder: Path) -> Iterable[Path]:
    """Yield valid .xlsx files and print skip reasons for others."""
    for path in sorted(folder.iterdir()):
        if not path.is_file():
            continue
        if path.name.startswith("~"):
            print(f"skipped file: {path.name} (temporary Excel lock file)")
            continue
        if path.suffix.lower() != ".xlsx":
            print(f"skipped file: {path.name} (not .xlsx)")
            continue
        yield path


def build_output_path(input_folder: Path, output_folder: Path) -> Path:
    """Create non-colliding output file path."""
    base = f"{input_folder.name}_PARAM"
    candidate = output_folder / f"{base}.xlsx"
    if not candidate.exists():
        return candidate

    idx = 1
    while True:
        candidate = output_folder / f"{base}.{idx}.xlsx"
        if not candidate.exists():
            return candidate
        idx += 1


def process_empirical_sheet(
    wb: xw.Book,
    metadata: Dict[str, str],
    source_file: str,
) -> List[Dict[str, Any]]:
    """Extract empirical candidates from one workbook."""
    try:
        sheet = wb.sheets["Empirical Model"]
    except Exception:
        print(f"skipped file: {source_file} (missing sheet 'Empirical Model')")
        return []

    anchor = find_anchor_cell(sheet, "max")
    if anchor is None:
        print(f"skipped file: {source_file} (no 'max' anchor on 'Empirical Model')")
        return []

    anchor_row, anchor_col = anchor
    n_quarters = 10

    # Offset map is anchored to the "max" cell layout.
    last_quarter_col = anchor_col - 10
    penetration_history_col = anchor_col - 9
    quarterly_sales_col = anchor_col - 7
    reported_sales_col = anchor_col - 6
    growth_rate_col = anchor_col - 5
    captured_in_db_col = anchor_col - 4
    forecast_col = anchor_col - 3
    actual_col = anchor_col - 2
    max_col = anchor_col
    min_col = anchor_col + 1

    helper_col = sheet.used_range.last_cell.column + 2
    helper_rows: List[Tuple[int, int]] = []

    for num_quarters in range(1, n_quarters + 1):
        if anchor_row - num_quarters < 1:
            continue
        helper_row = anchor_row + num_quarters
        avg_formula = (
            f"=AVERAGE(R{anchor_row - num_quarters}C{penetration_history_col}:"
            f"R{anchor_row - 1}C{penetration_history_col})"
        )
        set_formula2_r1c1(sheet.cells(helper_row, helper_col), avg_formula)
        helper_rows.append((num_quarters, helper_row))

    if helper_rows:
        wb.app.calculate()

    extracted_rows: List[Dict[str, Any]] = []
    for num_quarters, helper_row in helper_rows:
        row = anchor_row + num_quarters

        avg_penetration = to_float(sheet.cells(helper_row, helper_col).value)
        quarterly_sales = to_float(sheet.cells(row, quarterly_sales_col).value)
        reported_sales = to_float(sheet.cells(row, reported_sales_col).value)
        growth_rate = to_float(sheet.cells(row, growth_rate_col).value)
        captured_pct = to_float(sheet.cells(row, captured_in_db_col).value)
        forecast_value = to_float(sheet.cells(row, forecast_col).value)
        actual_value = to_float(sheet.cells(row, actual_col).value)
        forecast_max = to_float(sheet.cells(row, max_col).value)
        forecast_min = to_float(sheet.cells(row, min_col).value)
        last_quarter_used = sheet.cells(anchor_row - 1, last_quarter_col).value if anchor_row > 1 else None

        if forecast_value is None and avg_penetration is not None and quarterly_sales is not None:
            forecast_value = avg_penetration * quarterly_sales
        if actual_value is None:
            actual_value = reported_sales
        if forecast_max is None and forecast_value is not None:
            forecast_max = forecast_value
        if forecast_min is None and forecast_value is not None:
            forecast_min = forecast_value

        if all(
            value is None
            for value in (
                avg_penetration,
                forecast_value,
                forecast_max,
                forecast_min,
                quarterly_sales,
                reported_sales,
            )
        ):
            continue

        range_width = (
            forecast_max - forecast_min
            if forecast_max is not None and forecast_min is not None
            else None
        )

        extracted_rows.append(
            {
                "model": metadata["model"],
                "ticker": metadata["ticker"],
                "model_period": metadata["model_period"],
                "model_date": metadata["model_date"],
                "method": "empirical",
                "parameter_name": "avg_penetration_pct",
                "parameter_value": avg_penetration,
                "num_quarters_used": num_quarters,
                "last_quarter_used": last_quarter_used,
                "forecast_value": forecast_value,
                "actual_value": actual_value,
                "forecast_max": forecast_max,
                "forecast_min": forecast_min,
                "range_width": range_width,
                "avg_penetration_pct": avg_penetration,
                "quarterly_sales": quarterly_sales,
                "reported_sales": reported_sales,
                "growth_rate_pct": growth_rate,
                "sales_captured_in_db_pct": captured_pct,
                "source_file": source_file,
            }
        )

    return extracted_rows


def process_regression_sheet(
    wb: xw.Book,
    metadata: Dict[str, str],
    source_file: str,
) -> List[Dict[str, Any]]:
    """Extract regression candidates from one workbook."""
    try:
        sheet = wb.sheets["Regression Model"]
    except Exception:
        print(f"skipped file: {source_file} (missing sheet 'Regression Model')")
        return []

    anchor = find_anchor_cell(sheet, "max")
    if anchor is None:
        print(f"skipped file: {source_file} (no 'max' anchor on 'Regression Model')")
        return []

    anchor_row, anchor_col = anchor
    y_col = anchor_col - 7
    x_col = anchor_col - 11
    helper_col = sheet.used_range.last_cell.column + 2
    actual_col = anchor_col - 2

    work_rows: List[Tuple[int, int, int, int]] = []
    for idx, num_quarters in enumerate(range(2, 11), start=1):
        start_row = anchor_row - num_quarters
        end_row = anchor_row - 1
        if start_row < 1:
            continue

        row = anchor_row + idx
        work_rows.append((num_quarters, row, start_row, end_row))

        intercept_cell = sheet.cells(row, helper_col)
        slope_cell = sheet.cells(row, helper_col + 1)
        forecast_cell = sheet.cells(row, helper_col + 2)
        max_cell = sheet.cells(row, helper_col + 3)
        min_cell = sheet.cells(row, helper_col + 4)

        set_formula2_r1c1(
            intercept_cell,
            f"=INTERCEPT(R{start_row}C{y_col}:R{end_row}C{y_col},R{start_row}C{x_col}:R{end_row}C{x_col})",
        )
        set_formula2_r1c1(
            slope_cell,
            f"=SLOPE(R{start_row}C{y_col}:R{end_row}C{y_col},R{start_row}C{x_col}:R{end_row}C{x_col})",
        )
        set_formula2_r1c1(
            forecast_cell,
            f"=RC[-2]+RC[-1]*(R{end_row}C{x_col}+1)",
        )
        set_formula2_r1c1(max_cell, f"=MAX(R{start_row}C{y_col}:R{end_row}C{y_col})")
        set_formula2_r1c1(min_cell, f"=MIN(R{start_row}C{y_col}:R{end_row}C{y_col})")

    if work_rows:
        wb.app.calculate()

    extracted_rows: List[Dict[str, Any]] = []
    previous_signature: Optional[Tuple[Optional[float], ...]] = None

    for num_quarters, row, _, _ in work_rows:
        intercept = to_float(sheet.cells(row, helper_col).value)
        slope = to_float(sheet.cells(row, helper_col + 1).value)
        forecast_value = to_float(sheet.cells(row, helper_col + 2).value)
        forecast_max = to_float(sheet.cells(row, helper_col + 3).value)
        forecast_min = to_float(sheet.cells(row, helper_col + 4).value)
        actual_value = to_float(sheet.cells(anchor_row, actual_col).value)

        signature = (
            round(intercept, 10) if intercept is not None else None,
            round(slope, 10) if slope is not None else None,
            round(forecast_value, 10) if forecast_value is not None else None,
            round(forecast_max, 10) if forecast_max is not None else None,
            round(forecast_min, 10) if forecast_min is not None else None,
        )
        if signature == previous_signature:
            continue
        previous_signature = signature

        if all(value is None for value in (intercept, slope, forecast_value, forecast_max, forecast_min)):
            continue

        range_width = (
            forecast_max - forecast_min
            if forecast_max is not None and forecast_min is not None
            else None
        )

        extracted_rows.append(
            {
                "model": metadata["model"],
                "ticker": metadata["ticker"],
                "model_period": metadata["model_period"],
                "model_date": metadata["model_date"],
                "method": "regression",
                "parameter_name": "num_quarters_used",
                "parameter_value": num_quarters,
                "num_quarters_used": num_quarters,
                "forecast_value": forecast_value,
                "actual_value": actual_value,
                "forecast_max": forecast_max,
                "forecast_min": forecast_min,
                "range_width": range_width,
                "intercept": intercept,
                "slope": slope,
                "source_file": source_file,
            }
        )

    return extracted_rows


def apply_sheet_format(ws: Any, columns: Sequence[str]) -> None:
    """Apply requested output formatting."""
    for header_cell in ws[1]:
        header_cell.font = Font(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx, header in enumerate(columns, start=1):
        max_len = len(header)
        for row_values in ws.iter_rows(
            min_row=2,
            min_col=col_idx,
            max_col=col_idx,
            values_only=True,
        ):
            value = row_values[0]
            if value is None:
                continue
            cell_text = str(value)
            if len(cell_text) > max_len:
                max_len = len(cell_text)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 42)


def write_output_workbook(
    output_path: Path,
    empirical_rows: Sequence[Dict[str, Any]],
    regression_rows: Sequence[Dict[str, Any]],
) -> None:
    """Write the output workbook with both required sheets."""
    out_wb = Workbook()
    out_wb.remove(out_wb.active)

    empirical_ws = out_wb.create_sheet("empirical_candidates")
    empirical_ws.append(EMPIRICAL_COLUMNS)
    for row in empirical_rows:
        empirical_ws.append([row.get(col) for col in EMPIRICAL_COLUMNS])
    apply_sheet_format(empirical_ws, EMPIRICAL_COLUMNS)

    regression_ws = out_wb.create_sheet("regression_candidates")
    regression_ws.append(REGRESSION_COLUMNS)
    for row in regression_rows:
        regression_ws.append([row.get(col) for col in REGRESSION_COLUMNS])
    apply_sheet_format(regression_ws, REGRESSION_COLUMNS)

    out_wb.save(output_path)


def main() -> None:
    input_path = Path(input_dir).expanduser().resolve()
    output_path_dir = Path(output_dir).expanduser().resolve()

    if not input_path.exists() or not input_path.is_dir():
        raise FileNotFoundError(f"input_dir does not exist or is not a folder: {input_path}")

    output_path_dir.mkdir(parents=True, exist_ok=True)
    final_output_path = build_output_path(input_path, output_path_dir)

    empirical_rows: List[Dict[str, Any]] = []
    regression_rows: List[Dict[str, Any]] = []
    processed_count = 0

    app = xw.App(visible=False, add_book=False)
    app.display_alerts = False
    app.screen_updating = False
    try:
        app.enable_events = False
    except Exception:
        pass

    try:
        try:
            app.calculation = "manual"
        except Exception:
            pass

        for file_path in iterate_source_files(input_path):
            wb: Optional[xw.Book] = None
            try:
                wb = app.books.open(str(file_path), update_links=False)
                metadata = parse_file_labels(file_path)

                empirical_rows.extend(process_empirical_sheet(wb, metadata, file_path.name))
                regression_rows.extend(process_regression_sheet(wb, metadata, file_path.name))

                processed_count += 1
                print(f"processed file: {file_path.name}")
            except Exception as exc:
                print(f"skipped file: {file_path.name} (processing error: {exc})")
            finally:
                safe_close_workbook(wb)
    finally:
        try:
            app.calculation = "automatic"
        except Exception:
            pass
        app.quit()

    write_output_workbook(final_output_path, empirical_rows, regression_rows)

    print(f"output path: {final_output_path}")
    print(f"number of files processed: {processed_count}")
    print(f"number of empirical rows: {len(empirical_rows)}")
    print(f"number of regression rows: {len(regression_rows)}")


if __name__ == "__main__":
    main()
